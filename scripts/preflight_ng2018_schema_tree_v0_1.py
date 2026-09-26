#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
PROBE_PATH=ROOT/"scripts"/"probe_ng2018_solanaceae_sources_v0_1.py"
spec=importlib.util.spec_from_file_location("source_probe",PROBE_PATH)
probe=importlib.util.module_from_spec(spec)
spec.loader.exec_module(probe)

HEADER_ROW=1
TREE_KEYWORDS=("mcc","maximum clade credibility","treeannotator","beast")


def recover_trait_source()->tuple[bytes,str]:
    urls=[]
    for u in [probe.ARTICLE,probe.ARTICLE+"?login=false",probe.ARTICLE+"?searchresult=1",probe.ARTICLE_MINIMAL]:
        r=probe.get(u,"text/html,*/*")
        if r["ok"]:
            urls.extend(probe.extract_supplement_urls(r["body"]))
    urls.extend(probe.static_supplement_candidates())
    for u in dict.fromkeys(urls):
        r=probe.get(u)
        if r["ok"] and probe.xlsxish(r["body"]):
            return r["body"],r["final_url"]
    raise RuntimeError("exact Table S1 source no longer recoverable")


def recover_tree_source()->tuple[bytes,str]:
    urls=probe.treebase_candidate_urls(probe.TREEBASE_ID)
    for summary_url in probe.treebase_summary_urls(probe.TREEBASE_ID):
        r=probe.get(summary_url,"text/html,*/*")
        if r["ok"]:
            urls.extend(probe.extract_treebase_download_urls(r["body"],r["final_url"]))
    for u in dict.fromkeys(urls):
        r=probe.get(u,"application/xml,application/nexml+xml,text/plain,*/*")
        if r["ok"] and probe.treebase_payload(r["body"]):
            return r["body"],r["final_url"]
    raise RuntimeError("TreeBASE S23063 source no longer recoverable")


def workbook_header_metadata(body:bytes)->dict:
    wb=load_workbook(io.BytesIO(body),read_only=True,data_only=False)
    sheets=[]
    for ws in wb.worksheets:
        header=[
            "" if c.value is None else str(c.value).strip()
            for c in ws[HEADER_ROW]
        ]
        sheets.append({
            "title":ws.title,
            "max_row":ws.max_row,
            "max_column":ws.max_column,
            "header_row":HEADER_ROW,
            "header":header,
        })
    return {"sheet_count":len(sheets),"sheets":sheets}


def localname(tag:str)->str:
    return tag.rsplit("}",1)[-1]


def tree_metadata(body:bytes)->dict:
    if not body.lstrip().startswith(b"<"):
        return {
            "format":"nexus",
            "tree_count":None,
            "trees":[],
            "selection_status":"HOLD_NEXUS_METADATA_PARSER_NOT_IMPLEMENTED",
            "selected_tree_id":None,
        }
    root=ET.fromstring(body)
    trees=[]
    for t in root.iter():
        if localname(t.tag)!="tree":
            continue
        nodes=[x for x in t.iter() if localname(x.tag)=="node"]
        edges=[x for x in t.iter() if localname(x.tag)=="edge"]
        sources={e.attrib.get("source") for e in edges if e.attrib.get("source")}
        node_ids={n.attrib.get("id") for n in nodes if n.attrib.get("id")}
        leaf_ids=node_ids-sources
        trees.append({
            "id":t.attrib.get("id"),
            "label":t.attrib.get("label"),
            "xsi_type":next((v for k,v in t.attrib.items() if k.endswith("type")),None),
            "node_count":len(nodes),
            "edge_count":len(edges),
            "terminal_node_count":len(leaf_ids),
        })

    keyword_hits=[]
    for x in trees:
        text=(" ".join(str(v or "") for v in (x["id"],x["label"],x["xsi_type"]))).lower()
        if any(k in text for k in TREE_KEYWORDS):
            keyword_hits.append(x)

    if len(trees)==1:
        selected=trees[0]
        selection_status="UNIQUE_TREE_SELECTED_PRE_TRAIT"
    elif len(keyword_hits)==1:
        selected=keyword_hits[0]
        selection_status="UNIQUE_PUBLISHED_MCC_KEYWORD_TREE_SELECTED_PRE_TRAIT"
    else:
        selected=None
        selection_status="HOLD_TREE_SELECTION_AMBIGUOUS_PRE_TRAIT"

    return {
        "format":"nexml",
        "tree_count":len(trees),
        "trees":trees,
        "keyword_match_count":len(keyword_hits),
        "selection_status":selection_status,
        "selected_tree_id":selected["id"] if selected else None,
        "selected_tree_terminal_node_count":selected["terminal_node_count"] if selected else None,
    }


def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--out",type=Path,required=True)
    a=ap.parse_args()
    a.out.parent.mkdir(parents=True,exist_ok=True)

    trait_body,trait_url=recover_trait_source()
    tree_body,tree_url=recover_tree_source()
    wbmeta=workbook_header_metadata(trait_body)
    tmeta=tree_metadata(tree_body)

    if wbmeta["sheet_count"]!=1:
        status="HOLD_NG2018_WORKBOOK_SHEET_SELECTION_PRE_TRAIT"
    elif tmeta["selected_tree_id"] is None:
        status="HOLD_NG2018_TREE_SELECTION_PRE_TRAIT"
    else:
        status="NG2018_SCHEMA_TREE_METADATA_PREFLIGHT_READY"

    out={
        "version":"v0.1",
        "status":status,
        "trait_source":{
            "url":trait_url,
            "bytes":len(trait_body),
            "sha256":hashlib.sha256(trait_body).hexdigest(),
            **wbmeta,
        },
        "tree_source":{
            "url":tree_url,
            "bytes":len(tree_body),
            "sha256":hashlib.sha256(tree_body).hexdigest(),
            **tmeta,
        },
        "trait_header_rows_opened":1,
        "trait_data_rows_opened":0,
        "trait_state_frequencies_computed":False,
        "tree_tip_labels_crosswalked_to_traits":False,
        "hidden_memory_auc_computed":False,
        "next_gate":(
            "FREEZE_HEADER_COLUMN_MAPPING_AND_IDENTIFIER_CROSSWALK"
            if status=="NG2018_SCHEMA_TREE_METADATA_PREFLIGHT_READY"
            else "STOP_HOLD"
        ),
        "paper1_science_changed":False,
        "el_v0_3_science_changed":False,
        "v0_8_promotion_state_changed":False,
    }
    a.out.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
    print(json.dumps(out,indent=2,sort_keys=True))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
