#!/usr/bin/env python3
"""Audit Iris Supplementary Table 1 row identity/coding before any phylogenetic signal computation."""
from __future__ import annotations

import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import requests
from openpyxl import load_workbook

URL = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7588356/supplementaryFiles"
UA = {"User-Agent": "chun-falsification-audit/1.0 (+https://github.com/zuizui0223/chun)"}
OUT = Path("analysis/_generated/iris_trait_row_audit_v1.json")


def s(x):
    return "" if x is None else re.sub(r"\s+", " ", str(x)).strip()


def main():
    r = requests.get(URL, headers=UA, timeout=90)
    r.raise_for_status()
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        blob = zf.read("Table_1.xlsx")
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    ws = wb["Source of data"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [s(x) for x in rows[0]]
    idx = {h: i for i, h in enumerate(headers) if h}
    required = {"Species", "Colour", "Pigment"}
    if not required <= set(idx):
        raise SystemExit(f"missing required headers: {sorted(required-set(idx))}; headers={headers}")

    records = []
    for excel_row, row in enumerate(rows[1:], start=2):
        species = s(row[idx["Species"]])
        colour = s(row[idx["Colour"]]).lower()
        pigment = s(row[idx["Pigment"]]).lower()
        if species or colour or pigment:
            records.append({"excel_row": excel_row, "species": species, "colour": colour, "pigment": pigment})

    nonempty_species = [x for x in records if x["species"]]
    by_species = defaultdict(list)
    for x in nonempty_species:
        by_species[x["species"]].append(x)

    duplicate_species = {k: v for k, v in by_species.items() if len(v) > 1}
    colour_counts = Counter(x["colour"] for x in nonempty_species)
    pigment_counts = Counter(x["pigment"] for x in nonempty_species)

    # Determine whether trailing workbook content is additional data or formatting/source blocks.
    row_windows = []
    for lo in range(2, ws.max_row + 1, 100):
        hi = min(ws.max_row, lo + 99)
        chunk = [x for x in records if lo <= x["excel_row"] <= hi]
        row_windows.append({
            "excel_rows": [lo, hi],
            "n_nonempty_any": len(chunk),
            "n_species": sum(bool(x["species"]) for x in chunk),
            "n_colour": sum(bool(x["colour"]) for x in chunk),
            "n_pigment": sum(bool(x["pigment"]) for x in chunk),
        })

    # Raw abbreviations sheet is source metadata, safe to record before endpoint computation.
    ab = wb["abreviations"]
    abbreviations = [[s(v) for v in row] for row in ab.iter_rows(values_only=True)]

    out = {
        "source_url": URL,
        "sheet": "Source of data",
        "max_row": ws.max_row,
        "max_column": ws.max_column,
        "headers": headers,
        "n_records_with_species_or_colour_or_pigment": len(records),
        "n_rows_with_species": len(nonempty_species),
        "n_unique_species_strings": len(by_species),
        "n_duplicate_species_strings": len(duplicate_species),
        "duplicate_species_examples": {
            k: [{"excel_row": x["excel_row"], "colour": x["colour"], "pigment": x["pigment"]} for x in v]
            for k, v in list(sorted(duplicate_species.items()))[:20]
        },
        "colour_counts": dict(sorted(colour_counts.items())),
        "pigment_counts": dict(sorted(pigment_counts.items())),
        "rows_with_species_but_no_colour": [x for x in nonempty_species if not x["colour"]][:30],
        "rows_with_species_but_no_pigment": [x for x in nonempty_species if not x["pigment"]][:30],
        "first_10_species_rows": nonempty_species[:10],
        "last_10_species_rows": nonempty_species[-10:],
        "row_windows": row_windows,
        "abbreviations_sheet": abbreviations,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("IRIS_TRAIT_ROW_AUDIT=" + json.dumps(out, ensure_ascii=False))


if __name__ == "__main__":
    main()
