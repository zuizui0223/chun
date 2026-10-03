from __future__ import annotations

import importlib.util
from datetime import datetime
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "build_paper1_ajb_docx_v1_0.py"
spec = importlib.util.spec_from_file_location("paper1_docx_v1_0", SCRIPT)
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)


def _write_docx(path: Path, when: datetime, zip_time: tuple[int, int, int, int, int, int]) -> None:
    doc = Document()
    doc.add_paragraph("deterministic package")
    doc.core_properties.created = when
    doc.core_properties.modified = when
    doc.save(path)

    temp = path.with_suffix(".tmp.docx")
    with ZipFile(path, "r") as src, ZipFile(temp, "w") as dst:
        for member in src.infolist():
            info = ZipInfo(member.filename, date_time=zip_time)
            info.compress_type = member.compress_type
            info.external_attr = member.external_attr
            data = src.read(member.filename)
            if member.compress_type == ZIP_DEFLATED:
                dst.writestr(info, data, compress_type=ZIP_DEFLATED, compresslevel=9)
            else:
                dst.writestr(info, data, compress_type=member.compress_type)
    temp.replace(path)


def test_docx_normalization_is_byte_reproducible(tmp_path):
    first = tmp_path / "first.docx"
    second = tmp_path / "second.docx"
    _write_docx(first, datetime(2025, 1, 2, 3, 4, 5), (2025, 1, 2, 3, 4, 4))
    _write_docx(second, datetime(2026, 6, 7, 8, 9, 10), (2026, 6, 7, 8, 9, 10))

    mod.normalize_docx_package(first)
    mod.normalize_docx_package(second)

    assert first.read_bytes() == second.read_bytes()
    with ZipFile(first) as archive:
        assert all(info.date_time == mod.FIXED_ZIP_TIME for info in archive.infolist())
    props = Document(first).core_properties
    assert props.created == mod.FIXED_CORE_TIME
    assert props.modified == mod.FIXED_CORE_TIME
