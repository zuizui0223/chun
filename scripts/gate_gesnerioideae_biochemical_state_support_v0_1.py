#!/usr/bin/env python3
from __future__ import annotations

import argparse
import importlib.util
import io
import json
import re
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
CROSSWALK_SCRIPT=ROOT/"scripts"/"preflight_gesnerioideae_biochemical_crosswalk_v0_1.py"
spec=importlib.util.spec_from_file_location("gcw",CROSSWALK_SCRIPT)
gcw=importlib.util.module_from_spec(spec)
spec.loader.exec_module(gcw)

SHEET="FINAL SAMPLE LIST"
HEADER_ROW=3
FINE_COLUMNS=[
  "Pelargonidin-3-rutinoside",
  "Pelargonidin-7-glucoside",
  "Pelargonidin-3-sambubioside",
  "Cyanidin-3- rutinoside",
  "Cyanidin-3- glucoside",
  "Peonidin-3-rutinoside (or isomer)",
  "Delphinidin-3-glucoside (or isomer)",
  "Delphinidin- rhamnose-glucose",
  "Malvidin-3- rutinoside",
  "Apigenidinin-5-glucoside",
  "Luteolidinin-5-glucoside",
  "Trihydroxymethoxyl flavylium-glucuronide",
  "Unknown glycosylated anthocyanin",
]
RARE_MIN=5
MIN_COMMON=20


def canonical_header(v)->str:
    return "" if v is None else re.sub(r"\s+"," ",str(v).strip())


def bit_from_cell(v)->int:
    if v is None or (isinstance(v,str) and not v.strip()):
        return 0
    if isinstance(v,bool):
        raise ValueError(f"boolean chemistry token not allowed: {v!r}")
    if isinstance(v,(int,float)):
        if v < 0:
            raise ValueError(f"negative chemistry concentration not allowed: {v!r}")
        return int(v>0)
    raise ValueError(f"nonnumeric chemistry token: {v!r}")


def fine_state(values)->str:
    return "".join(str(bit_from_cell(v)) for v in values)


def coarse_state(fine:str)->str:
    return "PRESENT" if "1" in fine else "NONE"


def retain_supported_states(rows:list[dict],rare_min:int=RARE_MIN)->tuple[list[dict],dict]:
    counts=Counter(r["fine_state"] for r in rows)
    kept={s for s,n in counts.items() if n>=rare_min}
    retained=[r for r in rows if r["fine_state"] in kept]
    retained_fine=Counter(r["fine_state"] for r in retained)
    retained_coarse=Counter(r["coarse_state"] for r in retained)
    return retained,{
        "pre_filter_fine_state_count":len(counts),
        "pre_filter_fine_frequencies":dict(sorted(counts.items())),
        "retained_fine_state_count":len(retained_fine),
        "retained_fine_frequencies":dict(sorted(retained_fine.items())),
        "retained_coarse_state_count":len(retained_coarse),
        "retained_coarse_frequencies":dict(sorted(retained_coarse.items())),
    }


def extract_first_rows(body:bytes,crosswalk:dict)->list[dict]:
    matches={m["source_key"]:m["tree_tip"] for m in crosswalk["matches"]}
    wb=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
    ws=wb[SHEET]
    header=[canonical_header(c.value) for c in ws[HEADER_ROW]]
    if len(set(x for x in header if x)) != len([x for x in header if x]):
        raise ValueError("duplicate canonicalized source headers")
    required=["Species",*FINE_COLUMNS]
    missing=[x for x in required if x not in header]
    if missing:
        raise ValueError(f"frozen chemistry headers missing after canonicalization: {missing}")
    col={x:header.index(x)+1 for x in required}

    first=OrderedDict()
    for r in range(HEADER_ROW+1,ws.max_row+1):
        raw=ws.cell(r,col["Species"]).value
        if raw is None or not str(raw).strip():
            continue
        key=gcw.species_key(str(raw))
        if key not in matches or key in first:
            continue
        vals=[ws.cell(r,col[x]).value for x in FINE_COLUMNS]
        fs=fine_state(vals)
        first[key]={
            "source_key":key,
            "tree_tip":matches[key],
            "fine_state":fs,
            "coarse_state":coarse_state(fs),
        }
    return list(first.values())


def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--crosswalk-json",type=Path,required=True)
    ap.add_argument("--out",type=Path,required=True)
    ap.add_argument("--frame-out",type=Path,required=True)
    a=ap.parse_args()
    a.out.parent.mkdir(parents=True,exist_ok=True)
    a.frame_out.parent.mkdir(parents=True,exist_ok=True)

    cw=json.loads(a.crosswalk_json.read_text())
    if cw["status"]!="GESNERIOIDEAE_BIOCHEMICAL_CROSSWALK_FROZEN_CHEMISTRY_UNOPENED":
        raise RuntimeError(f"crosswalk not ready: {cw['status']}")

    body=gcw.exact_workbook()
    rows=extract_first_rows(body,cw)
    retained,stats=retain_supported_states(rows)

    common=len(retained)
    fine_n=stats["retained_fine_state_count"]
    coarse_n=stats["retained_coarse_state_count"]

    if common<MIN_COMMON:
        status="HOLD_COMMON_FRAME_LT_20"
    elif fine_n<2:
        status="HOLD_FINE_STATE_SUPPORT_LT_2"
    elif coarse_n<2:
        status="HOLD_COARSE_STATE_SUPPORT_LT_2"
    elif fine_n<=coarse_n:
        status="STRUCTURAL_NO_HIDDEN_MEMORY_TEST"
    else:
        status="GESNERIOIDEAE_BIOCHEMICAL_HIDDEN_MEMORY_OPPORTUNITY_CONFIRMED"

    frame={
      "version":"v0.1",
      "rows":retained,
      "retained_tips":[r["tree_tip"] for r in retained],
      "chemistry_values_opened":True,
      "hidden_memory_auc_computed":False,
    }
    a.frame_out.write_text(json.dumps(frame,indent=2,sort_keys=True)+"\n")

    out={
      "version":"v0.1",
      "status":status,
      "source_crosswalk_version":cw.get("version"),
      "header_normalization":"strip + collapse Unicode/ASCII whitespace to one ASCII space, identical to pre-outcome header probe",
      "frozen_compound_columns":FINE_COLUMNS,
      "detection_rule":"numeric >0 => 1; numeric 0 or blank => 0; any other nonblank/nonnumeric token => HOLD",
      "rare_fine_state_minimum_tips":RARE_MIN,
      "crosswalk_matched_species":cw["matched_species"],
      "chemistry_rows_read":len(rows),
      "common_frame_tips":common,
      **stats,
      "compression_opportunity":bool(fine_n>coarse_n),
      "chemistry_values_opened":True,
      "state_frequencies_computed":True,
      "hidden_memory_auc_computed":False,
      "next_gate":("RUN_FROZEN_WITHIN_COARSE_HIDDEN_MEMORY_TEST"
                   if status=="GESNERIOIDEAE_BIOCHEMICAL_HIDDEN_MEMORY_OPPORTUNITY_CONFIRMED"
                   else "STOP"),
      "paper1_science_changed":False,
      "el_v0_3_science_changed":False,
    }
    a.out.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
    print(json.dumps(out,indent=2,sort_keys=True))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
