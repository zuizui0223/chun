#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT=Path(__file__).resolve().parents[1]
MAN=ROOT/"manuscript"/"CROSS_RADIATION_EVOLUTION_LETTERS_V0_3_TEMPORAL_MEMORY_CANDIDATE.md"
META=ROOT/"submission"/"CROSS_RADIATION_EL_SUBMISSION_METADATA_V0_2.json"

HEADING_RE=re.compile(r"^(#{1,3})\s+(.*)$")
NUMBERED_RE=re.compile(r"^\d+\.\s+(.*)$")
INLINE_RE=re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|\x60[^\x60]+\x60)")


def add_runs(p,text:str):
    pos=0
    for m in INLINE_RE.finditer(text):
        if m.start()>pos:
            p.add_run(text[pos:m.start()])
        tok=m.group(0)
        if tok.startswith("**"):
            r=p.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith("*"):
            r=p.add_run(tok[1:-1]); r.italic=True
        elif tok.startswith(chr(96)):
            r=p.add_run(tok[1:-1]); r.font.name="Courier New"
        pos=m.end()
    if pos<len(text):
        p.add_run(text[pos:])


def replace_section(md:str,heading:str,next_heading:str|None,body:str)->str:
    start=md.find(heading)
    if start<0:
        return md
    content_start=start+len(heading)
    end=md.find(next_heading,content_start) if next_heading else len(md)
    if end<0:
        end=len(md)
    return md[:content_start]+"\n\n"+body.strip()+"\n\n"+md[end:]


def humanize(md:str,meta:dict)->tuple[str,list[str]]:
    p1=meta["phase1_author_identity"]
    p2=meta["phase2_declarations"]
    p3=meta["phase3_archive"]
    authors=p1.get("authors") or []
    affs=p1.get("affiliations") or []
    corr=p1.get("corresponding_author") or {}
    notes=[]

    if authors and affs:
        ordered=[]
        for a in authors:
            ids=",".join(a.get("affiliation_ids") or [])
            ordered.append(f"{a['name']} [{ids}]" if ids else a["name"])
        aff_lines=[f"[{a['id']}] {a['institution']}, {a['address']}" for a in affs]
        author_block="Authors: "+"; ".join(ordered)+"\n\nAffiliations:\n" + "\n".join("- "+x for x in aff_lines)
        if corr.get("author_name") and corr.get("email"):
            author_block += f"\n\nCorresponding author: {corr['author_name']}; {corr['email']}; {corr.get('postal_address') or ''}"
        marker="Running title:"
        i=md.find(marker)
        if i>=0:
            j=md.find("\n",i)
            if j>=0:
                md=md[:j+1]+"\n"+author_block+"\n"+md[j+1:]
    else:
        notes.append("author identity pending")

    credits=p2.get("credit_roles") or []
    if credits:
        body="\n".join(f"{r['author_name']}: {', '.join(r['roles'])}." for r in credits)
        md=replace_section(md,"# Author contributions","# Funding",body)
    else:
        notes.append("CRediT pending")

    status=p2.get("funding_status")
    if status=="none":
        md=replace_section(md,"# Funding","# Conflict of interest","No external funding was received for this work.")
    elif status=="funded" and p2.get("funding_sources"):
        md=replace_section(md,"# Funding","# Conflict of interest"," ".join(str(x).strip() for x in p2["funding_sources"] if str(x).strip()))
    else:
        notes.append("funding pending")

    ack=p2.get("acknowledgments")
    if isinstance(ack,str) and ack.strip():
        text=ack.strip()
        if text.lower()=="none":
            text="None."
        md=replace_section(md,"# Acknowledgements","# Figure legends and alt text",text)
    else:
        notes.append("acknowledgements pending")

    doi=p3.get("archive_doi")
    url=p3.get("archive_url")
    version=p3.get("archive_version")
    if doi and url and version:
        old="A submission-specific CHUN code snapshot will be archived with a persistent identifier before submission."
        new=f"The submission-specific CHUN code snapshot is archived as version {version} at {url} (DOI: {doi})."
        md=md.replace(old,new)
    else:
        notes.append("archive DOI pending")
    return md,notes


def set_line_numbering(doc:Document):
    for section in doc.sections:
        sectPr=section._sectPr
        for child in list(sectPr):
            if child.tag==qn("w:lnNumType"):
                sectPr.remove(child)
        ln=OxmlElement("w:lnNumType")
        ln.set(qn("w:countBy"),"1")
        ln.set(qn("w:start"),"1")
        ln.set(qn("w:restart"),"continuous")
        sectPr.append(ln)


def style_document(doc:Document):
    for name in ("Normal","Heading 1","Heading 2","Heading 3","List Bullet","List Number"):
        if name not in doc.styles:
            continue
        st=doc.styles[name]
        st.font.name="Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"),"Times New Roman")
        st.font.size=Pt(12)
        st.paragraph_format.line_spacing_rule=WD_LINE_SPACING.DOUBLE
        st.paragraph_format.space_after=Pt(0)
        st.paragraph_format.space_before=Pt(0)
    for section in doc.sections:
        section.top_margin=Inches(1)
        section.bottom_margin=Inches(1)
        section.left_margin=Inches(1)
        section.right_margin=Inches(1)


def add_page_number(doc:Document):
    for section in doc.sections:
        p=section.footer.paragraphs[0]
        p.alignment=2
        fld=OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"),"PAGE")
        p._p.append(fld)


def render_markdown(doc:Document,md:str):
    for raw in md.splitlines():
        line=raw.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        m=HEADING_RE.match(line)
        if m:
            level=len(m.group(1))
            p=doc.add_paragraph(style=f"Heading {level}")
            add_runs(p,m.group(2))
            continue
        if line.startswith("- "):
            p=doc.add_paragraph(style="List Bullet")
            add_runs(p,line[2:])
            continue
        m=NUMBERED_RE.match(line)
        if m:
            p=doc.add_paragraph(style="List Number")
            add_runs(p,m.group(1))
            continue
        if line.startswith("> "):
            p=doc.add_paragraph()
            p.paragraph_format.left_indent=Inches(0.25)
            add_runs(p,line[2:])
            continue
        p=doc.add_paragraph()
        add_runs(p,line)


def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--out",type=Path,required=True)
    ap.add_argument("--status",type=Path,required=True)
    a=ap.parse_args()
    a.out.parent.mkdir(parents=True,exist_ok=True)
    a.status.parent.mkdir(parents=True,exist_ok=True)

    md=MAN.read_text(encoding="utf-8")
    meta=json.loads(META.read_text(encoding="utf-8"))
    md,notes=humanize(md,meta)

    doc=Document()
    style_document(doc)
    set_line_numbering(doc)
    add_page_number(doc)
    render_markdown(doc,md)
    props=doc.core_properties
    props.title=meta["paper_title"]
    props.subject="Evolution Letters v0.3 submission manuscript"
    props.comments="Generated from frozen v0.3 science source; human metadata may remain pending."
    doc.save(a.out)

    check=Document(a.out)
    if len(check.paragraphs)<100:
        raise RuntimeError("unexpectedly short DOCX")
    if "lnNumType" not in check.sections[0]._sectPr.xml:
        raise RuntimeError("line numbering missing")
    out={
      "version":"v0.1",
      "status":"EL_V0_3_PREHUMAN_DOCX_READY" if notes else "EL_V0_3_FINAL_DOCX_CONTENT_READY",
      "docx":str(a.out),
      "paragraphs":len(check.paragraphs),
      "line_numbering":"continuous",
      "font":"Times New Roman 12 pt",
      "spacing":"double",
      "page_numbers":True,
      "human_metadata_pending":notes,
      "source_manuscript":str(MAN.relative_to(ROOT)),
      "source_manuscript_modified":False,
      "scientific_results_changed":False
    }
    a.status.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
    print(json.dumps(out,indent=2,sort_keys=True))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
