#!/usr/bin/env python3
"""Retrieve and inspect the frozen Nicotiana supplementary source without computing phylogenetic signal."""
from __future__ import annotations

import hashlib
import io
import json
import re
import subprocess
import zipfile
from pathlib import Path

import requests
from docx import Document

PMC_ID = "PMC4598364"
SUPPLEMENT = f"https://www.ebi.ac.uk/europepmc/webservices/rest/{PMC_ID}/supplementaryFiles"
OUT = Path("analysis/_generated/nicotiana_source_v1")
UA = {"User-Agent": "chun-nicotiana-falsification/1.0 (+https://github.com/zuizui0223/chun)"}


def clean(x):
    return re.sub(r"\s+", " ", str(x or "")).strip()


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    r = requests.get(SUPPLEMENT, headers=UA, timeout=120)
    r.raise_for_status()
    package_sha = hashlib.sha256(r.content).hexdigest()
    if r.content[:2] != b"PK":
        raise SystemExit("Europe PMC supplementaryFiles response is not a ZIP")

    inventory = {
        "source_doi": "10.1093/aob/mcv048",
        "pmc_id": PMC_ID,
        "supplementary_package_url": SUPPLEMENT,
        "package_bytes": len(r.content),
        "package_sha256": package_sha,
        "members": [],
        "selected_doc": None,
        "converted_docx": None,
        "tables": [],
        "paragraph_hits": [],
    }

    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        members = zf.infolist()
        for m in members:
            inventory["members"].append({"name": m.filename, "bytes": m.file_size})
        docs = [m for m in members if m.filename.lower().endswith(".doc")]
        if len(docs) != 1:
            raise SystemExit(f"expected one legacy Word supplement, found {[m.filename for m in docs]}")
        m = docs[0]
        raw = zf.read(m)
        raw_path = OUT / "source_supplement.doc"
        raw_path.write_bytes(raw)
        inventory["selected_doc"] = {
            "member": m.filename,
            "bytes": len(raw),
            "sha256": hashlib.sha256(raw).hexdigest(),
            "ole_magic": raw[:8].hex(),
        }

    # Conversion is a source-format operation only. No trait selection occurs here.
    proc = subprocess.run(
        ["libreoffice", "--headless", "--convert-to", "docx", "--outdir", str(OUT), str(raw_path)],
        text=True, capture_output=True, timeout=180,
    )
    inventory["libreoffice"] = {
        "returncode": proc.returncode,
        "stdout": clean(proc.stdout),
        "stderr": clean(proc.stderr),
    }
    if proc.returncode != 0:
        raise SystemExit(f"LibreOffice conversion failed: {proc.stderr}")

    candidates = [p for p in OUT.glob("*.docx")]
    if len(candidates) != 1:
        raise SystemExit(f"expected exactly one converted docx, found {[p.name for p in candidates]}")
    docx_path = candidates[0]
    inventory["converted_docx"] = {
        "path": docx_path.name,
        "bytes": docx_path.stat().st_size,
        "sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest(),
    }

    doc = Document(docx_path)
    keywords = ("table s1", "table s2", "table s4", "diploid", "polyploid", "hybrid", "chlorophyll", "spectral")
    for i, p in enumerate(doc.paragraphs):
        t = clean(p.text)
        if t and any(k in t.lower() for k in keywords):
            inventory["paragraph_hits"].append({"paragraph_index": i, "text": t})

    for ti, table in enumerate(doc.tables):
        rows = []
        for ri, row in enumerate(table.rows):
            vals = [clean(cell.text) for cell in row.cells]
            rows.append(vals)
        max_cols = max((len(x) for x in rows), default=0)
        nonempty = sum(any(v for v in row) for row in rows)
        preview = rows[:8]
        flattened = " ".join(v for row in rows for v in row).lower()
        tags = [k.upper().replace(" ", "_") for k in keywords if k in flattened]
        tsv = OUT / f"table_{ti:02d}.tsv"
        with tsv.open("w", encoding="utf-8") as f:
            for row in rows:
                f.write("\t".join(v.replace("\t", " ").replace("\n", " ") for v in row) + "\n")
        inventory["tables"].append({
            "table_index": ti,
            "rows": len(rows),
            "nonempty_rows": nonempty,
            "max_columns": max_cols,
            "tags": tags,
            "preview_first_8_rows": preview,
            "tsv": tsv.name,
        })

    # Source-schema matching only: identify likely S1/S2 by explicit table labels/column semantics.
    s1_candidates = []
    s2_candidates = []
    for t in inventory["tables"]:
        blob = " ".join(v for row in t["preview_first_8_rows"] for v in row).lower()
        if "accession" in blob and ("tax" in blob or "species" in blob):
            s1_candidates.append(t["table_index"])
        if ("spectral" in blob or "colour" in blob or "color" in blob) and "chlorophyll" in blob:
            s2_candidates.append(t["table_index"])
    inventory["schema_candidates"] = {"S1": s1_candidates, "S2": s2_candidates}
    inventory["schema_status"] = "UNIQUE" if len(s1_candidates) == 1 and len(s2_candidates) == 1 else "NEEDS_AUDIT"

    (OUT / "inventory.json").write_text(json.dumps(inventory, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    compact = {
        "package_sha256": package_sha,
        "selected_doc": inventory["selected_doc"],
        "converted_docx": inventory["converted_docx"],
        "n_tables": len(inventory["tables"]),
        "tables": inventory["tables"],
        "schema_candidates": inventory["schema_candidates"],
        "schema_status": inventory["schema_status"],
        "paragraph_hits": inventory["paragraph_hits"][:30],
    }
    print("NICOTIANA_SOURCE_INSPECTION=" + json.dumps(compact, ensure_ascii=False))


if __name__ == "__main__":
    main()
