#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import urllib.request
from pathlib import Path

from Bio import Phylo
from openpyxl import load_workbook

ARTICLE_DOI="10.1007/s10681-026-03700-1"
SUPP_URL="https://media.springernature.com/original/springer-static/esm/art%3A10.1007%2Fs10681-026-03700-1/MediaObjects/10681_2026_3700_MOESM1_ESM.xlsx"
AUTHOR_REPO="samehrem/Mehrem_etal_2025_PhenoLac"
AUTHOR_COMMIT="1abe8019b9c90ba102990dcd5f2f97cadc85d925"
AUTHOR_SCRIPT="Figure_3_S3.R"
AUTHOR_SCRIPT_BLOB="d722a3bbe9e4d70beddd6b3ce64a3d03573ce9a5"
RAW_SCRIPT=f"https://raw.githubusercontent.com/{AUTHOR_REPO}/{AUTHOR_COMMIT}/{AUTHOR_SCRIPT}"
UA="CHUN-Lactuca-source-first/0.1"
EXPECTED_MAPPING={
  "0":"white",
  "1":"light yellow",
  "2":"yellow",
  "3":"strong yellow",
  "4":"light purple",
  "5":"purple",
  "6":"pink",
}

def fetch(url:str,accept:str="*/*")->bytes:
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":accept})
    with urllib.request.urlopen(req,timeout=120) as r:
        return r.read()

def sha256(b:bytes)->str:
    return hashlib.sha256(b).hexdigest()

def normalize_header(x)->str:
    return re.sub(r"\s+"," ",str(x or "").strip())

def script_contract(body:bytes)->dict:
    txt=body.decode("utf-8",errors="replace")
    hits=[]
    for i,line in enumerate(txt.splitlines(),1):
        if (
          "Supplemental_table_LactucaPhenotypes_2025.xlsx" in line
          or "PetColor" in line
          or "0 = white" in line
          or 'sheet="Lactuca_tree"' in line
          or 'sheet="S2_Phenotypes"' in line
        ):
            hits.append({"line":i,"text":line})
    mapping_line=next((x["text"] for x in hits if "0 = white" in x["text"]),None)
    mapping_ok=bool(mapping_line and all(v in mapping_line.lower() for v in EXPECTED_MAPPING.values()))
    return {"sha256":sha256(body),"evidence":hits,"mapping_line":mapping_line,"mapping_ok":mapping_ok}

def workbook_metadata(body:bytes)->dict:
    wb=load_workbook(io.BytesIO(body),read_only=True,data_only=False)
    sheets=wb.sheetnames
    out={"sheets":sheets}
    if "S2_Phenotypes" not in sheets or "Lactuca_tree" not in sheets:
        out["required_sheets_ready"]=False
        return out
    ws=wb["S2_Phenotypes"]
    header=[normalize_header(c.value) for c in next(ws.iter_rows(min_row=1,max_row=1))]
    out["required_sheets_ready"]=True
    out["phenotype_sheet_max_row"]=ws.max_row
    out["phenotype_sheet_max_column"]=ws.max_column
    out["phenotype_header"]=header
    out["species_header_present"]="Species" in header
    out["petcolor_header_present"]="PetColor" in header

    tws=wb["Lactuca_tree"]
    tree_text=str(tws.cell(1,1).value or "").strip()
    out["tree_cell_nonempty"]=bool(tree_text)
    out["tree_text_sha256"]=sha256(tree_text.encode("utf-8"))
    if tree_text:
        tree=Phylo.read(io.StringIO(tree_text),"newick")
        tips=tree.get_terminals()
        branches=[c.branch_length for c in tree.find_clades() if c is not tree.root]
        out["tree_tip_count"]=len(tips)
        out["tree_duplicate_tip_labels"]=len(tips)-len({str(t.name or "").strip() for t in tips})
        out["tree_missing_nonroot_branch_lengths"]=sum(x is None for x in branches)
        out["tree_negative_nonroot_branch_lengths"]=sum(x is not None and x<0 for x in branches)
    return out

def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--out",type=Path,required=True)
    a=ap.parse_args(); a.out.parent.mkdir(parents=True,exist_ok=True)
    try:
        workbook=fetch(SUPP_URL)
        script=fetch(RAW_SCRIPT,"text/plain,*/*")
        sc=script_contract(script)
        wm=workbook_metadata(workbook)
        exact_ready=bool(
          workbook.startswith(b"PK\x03\x04")
          and sc["mapping_ok"]
          and wm.get("required_sheets_ready")
          and wm.get("species_header_present")
          and wm.get("petcolor_header_present")
          and wm.get("tree_tip_count",0)>=20
          and wm.get("tree_duplicate_tip_labels")==0
          and wm.get("tree_missing_nonroot_branch_lengths")==0
          and wm.get("tree_negative_nonroot_branch_lengths")==0
        )
        status="PASS_LACTUCA_SOURCE_FIRST_EXACT_WORKBOOK_TREE_AND_SCHEMA_READY" if exact_ready else "HOLD_LACTUCA_SOURCE_OR_SCHEMA_INCOMPLETE"
        out={
          "version":"v0.1",
          "status":status,
          "candidate":"LACTUCA",
          "candidate_order":4,
          "article_doi":ARTICLE_DOI,
          "official_supplement":{
            "url":SUPP_URL,
            "bytes":len(workbook),
            "sha256":sha256(workbook),
            "xlsx_signature":workbook.startswith(b"PK\x03\x04"),
          },
          "author_code":{
            "repo":AUTHOR_REPO,
            "commit":AUTHOR_COMMIT,
            "path":AUTHOR_SCRIPT,
            "git_blob_sha":AUTHOR_SCRIPT_BLOB,
            **sc,
          },
          "pre_row_colour_schema":{
            "raw_code_mapping":EXPECTED_MAPPING,
            "fine_states":list(EXPECTED_MAPPING.values()),
            "coarse_state":"WHITE iff code 0/source label white; codes 1-6 NONWHITE",
            "mapping_frozen_before_species_level_PetColor_values":True,
          },
          "workbook_metadata":wm,
          "hard_criteria":{
            "reported_accessions":550,
            "reported_wild_relatives":20,
            "tree_tips_ge_20":wm.get("tree_tip_count",0)>=20,
            "literal_or_author_mapped_discrete_colour":sc["mapping_ok"],
            "exact_trait_and_tree_same_official_object":True,
            "branch_lengths_present":wm.get("tree_missing_nonroot_branch_lengths")==0,
          },
          "outcome_firewall":{
            "phenotype_header_opened":True,
            "species_level_PetColor_values_opened":False,
            "colour_state_frequencies_computed":False,
            "tree_tip_labels_emitted":False,
            "hidden_memory_auc_computed":False,
          },
          "next_gate":"STOP_SCREEN_AND_ADMIT_LACTUCA_IDENTIFIER_ONLY_CROSSWALK" if exact_ready else "CONTINUE_SOURCE_FIRST_SCREEN_IF_BUDGET_REMAINS",
          "paper1_science_changed":False,
          "el_v0_3_science_changed":False
        }
    except Exception as e:
        out={
          "version":"v0.1","status":"HOLD_LACTUCA_OFFICIAL_SUPPLEMENT_TRANSPORT_OR_PARSE",
          "candidate":"LACTUCA","candidate_order":4,
          "error":f"{type(e).__name__}: {e}",
          "outcome_firewall":{"species_level_PetColor_values_opened":False,"colour_state_frequencies_computed":False,"tree_tip_labels_emitted":False,"hidden_memory_auc_computed":False},
          "next_gate":"CONTINUE_SOURCE_FIRST_SCREEN_IF_BUDGET_REMAINS",
          "paper1_science_changed":False,"el_v0_3_science_changed":False
        }
    a.out.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
    print(json.dumps({
      "status":out["status"],
      "supplement":out.get("official_supplement"),
      "author_mapping":(out.get("author_code") or {}).get("mapping_line"),
      "workbook_metadata":out.get("workbook_metadata"),
      "next_gate":out["next_gate"],
    },indent=2))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
