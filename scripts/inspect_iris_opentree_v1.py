#!/usr/bin/env python3
"""Audit the frozen Iris OpenTree matching/admission layer without computing colour signal."""
from __future__ import annotations

import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import requests
from Bio import Phylo
from openpyxl import load_workbook

SUPPLEMENT = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7588356/supplementaryFiles"
OT_API = "https://api.opentreeoflife.org/v3"
UA = "chun-iris-falsification/1.0"
OUT_DIR = Path("analysis/_generated/iris_opentree_v1")
RANKS = {"subsp.", "ssp.", "var.", "f."}


def get_species_rows():
    r = requests.get(SUPPLEMENT, headers={"User-Agent": UA}, timeout=90)
    r.raise_for_status()
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        wb = load_workbook(io.BytesIO(zf.read("Table_1.xlsx")), read_only=True, data_only=True)
    ws = wb["Source of data"]
    rows = list(ws.iter_rows(values_only=True))
    headers = ["" if x is None else str(x).strip() for x in rows[0]]
    i = headers.index("Species")
    species = [re.sub(r"\s+", " ", str(row[i])).strip() for row in rows[1:] if row[i] not in (None, "")]
    if len(species) != 226 or len(set(species)) != 226:
        raise SystemExit(f"frozen source row identity changed: n={len(species)} unique={len(set(species))}")
    return species


def canonical_query(source_name: str) -> tuple[str, str]:
    """Return colour-blind TNRS query and deterministic source-format normalization tags."""
    raw = re.sub(r"\s+", " ", source_name.strip())
    tags: list[str] = []
    epithet = None
    rest = ""

    # Attached publisher hybrid notation, e.g. `Irisxgermanica L.`.
    m = re.match(r"(?i)^iris[x×]([a-z][a-z-]*)(?:\s+(.*))?$", raw)
    if m:
        epithet = m.group(1)
        rest = m.group(2) or ""
        tags.append("ATTACHED_HYBRID_MARKER_REMOVED")
    else:
        # Separated hybrid notation, e.g. `Iris x germanica ...`.
        m = re.match(r"(?i)^iris\s+[x×]\s+([a-z][a-z-]*)(?:\s+(.*))?$", raw)
        if m:
            epithet = m.group(1)
            rest = m.group(2) or ""
            tags.append("SEPARATED_HYBRID_MARKER_REMOVED")
        else:
            # Ordinary binomial. This deliberately requires whitespace so an
            # epithet beginning with x (e.g. Iris xiphium) is not misread as a hybrid marker.
            m = re.match(r"(?i)^iris\s+([a-z][a-z-]*)(?:\s+(.*))?$", raw)
            if not m:
                raise ValueError("NO_IRIS_BINOMIAL_PATTERN")
            epithet = m.group(1)
            rest = m.group(2) or ""

    if not raw.startswith("Iris"):
        tags.append("GENUS_CASE_NORMALIZED")

    base = ["Iris", epithet.lower()]
    tail = rest.split()
    if len(tail) >= 2 and tail[0].lower() in RANKS:
        base.extend([tail[0].lower(), tail[1].lower()])
    return " ".join(base), "+".join(tags) if tags else "NONE"


def post(path, payload):
    r = requests.post(f"{OT_API}/{path}", json=payload, headers={"User-Agent": UA}, timeout=120)
    r.raise_for_status()
    return r.json()


def tnrs(names):
    data = post("tnrs/match_names", {
        "names": names,
        "do_approximate_matching": False,
        "include_suppressed": False,
    })
    results = data.get("results", [])
    if len(results) != len(names):
        raise SystemExit(f"TNRS returned {len(results)} results for {len(names)} names")
    out = []
    for query, result in zip(names, results):
        matches = result.get("matches", [])
        good = [m for m in matches if float(m.get("score", 1)) >= 0.999999 and not m.get("is_approximate_match", False)]
        if len(good) == 1:
            m = good[0]
            tx = m.get("taxon", {})
            out.append({
                "query": query,
                "ott_id": tx.get("ott_id"),
                "matched_name": tx.get("name", ""),
                "is_synonym": bool(m.get("is_synonym", False)),
                "score": float(m.get("score", 1)),
                "status": "EXACT",
            })
        else:
            out.append({
                "query": query,
                "ott_id": None,
                "matched_name": "",
                "is_synonym": None,
                "score": None,
                "status": f"REJECT_{len(good)}_EXACT_MATCHES",
            })
    return out


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    source_names = get_species_rows()

    parsed = []
    parse_errors = []
    for source in source_names:
        try:
            query, normalization = canonical_query(source)
            parsed.append((source, query, normalization))
        except Exception as e:
            parse_errors.append({"source_species": source, "error": f"{type(e).__name__}: {e}"})

    parser_summary = {
        "n_source_rows": len(source_names),
        "n_parsed": len(parsed),
        "n_parse_errors": len(parse_errors),
        "parse_errors": parse_errors,
        "normalization_counts": dict(Counter(x[2] for x in parsed)),
    }
    (OUT_DIR / "source_name_parser_audit.json").write_text(
        json.dumps(parser_summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    if parse_errors:
        print("IRIS_SOURCE_NAME_PARSER_AUDIT=" + json.dumps(parser_summary, ensure_ascii=False))
        raise SystemExit(f"source-name parser unresolved for {len(parse_errors)} rows; TNRS not opened")

    canonical = [x[1] for x in parsed]
    normalizations = [x[2] for x in parsed]
    canonical_dups = {q: n for q, n in Counter(canonical).items() if n > 1}

    matches = tnrs(canonical)
    rows = []
    by_ott = defaultdict(list)
    for (source, query, normalization), m in zip(parsed, matches):
        row = {
            "source_species": source,
            "query": query,
            "source_format_normalization": normalization,
            **{k: v for k, v in m.items() if k != "query"},
        }
        rows.append(row)
        if row["status"] == "EXACT" and row["ott_id"] is not None:
            by_ott[int(row["ott_id"])].append(row)

    collisions = {oid: rs for oid, rs in by_ott.items() if len(rs) > 1}
    collision_sources = {r["source_species"] for rs in collisions.values() for r in rs}
    for row in rows:
        if row["source_species"] in collision_sources:
            row["status"] = "REJECT_OTT_COLLISION"

    admitted = [r for r in rows if r["status"] == "EXACT"]
    ott_to_source = {int(r["ott_id"]): r["source_species"] for r in admitted}
    ids = sorted(ott_to_source)
    if not ids:
        raise SystemExit("no admitted OpenTree ids")

    sub = post("tree_of_life/induced_subtree", {"ott_ids": ids, "label_format": "id"})
    nwk = sub.get("newick", "")
    if not nwk:
        raise SystemExit("OpenTree returned empty induced subtree")
    (OUT_DIR / "opentree_induced_raw.nwk").write_text(nwk + "\n", encoding="utf-8")

    tree = Phylo.read(io.StringIO(nwk), "newick")
    tree_ott = []
    for tip in tree.get_terminals():
        label = str(tip.name).strip("'")
        m = re.search(r"(?:ott)?(\d+)$", label)
        if m:
            tree_ott.append(int(m.group(1)))
    tree_ott_set = set(tree_ott)
    admitted_set = set(ids)
    missing_from_tree = sorted(admitted_set - tree_ott_set)
    unexpected_tree_ids = sorted(tree_ott_set - admitted_set)
    tree_overlap = len(admitted_set & tree_ott_set)

    summary = {
        "n_source_rows": len(source_names),
        "n_unique_source_rows": len(set(source_names)),
        "n_unique_canonical_queries": len(set(canonical)),
        "source_format_normalization_counts": dict(Counter(normalizations)),
        "canonical_query_collisions": canonical_dups,
        "tnrs_status_counts": dict(Counter(r["status"] for r in rows)),
        "n_ott_collision_groups": len(collisions),
        "ott_collision_groups": {
            str(oid): [{"source_species": r["source_species"], "query": r["query"], "matched_name": r["matched_name"]} for r in rs]
            for oid, rs in collisions.items()
        },
        "n_admitted_unique_ott": len(ids),
        "n_tree_overlap": tree_overlap,
        "coverage_fraction_of_226": tree_overlap / 226.0,
        "admission_gate_80pct": tree_overlap >= 181,
        "missing_admitted_ids_from_tree": missing_from_tree,
        "unexpected_tree_ids": unexpected_tree_ids,
        "synth_id": sub.get("synth_id", ""),
    }
    (OUT_DIR / "tnrs_matches.json").write_text(json.dumps(rows, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    (OUT_DIR / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("IRIS_OPENTREE_AUDIT=" + json.dumps(summary, ensure_ascii=False))
    if not summary["admission_gate_80pct"]:
        raise SystemExit("OpenTree admission gate failed: fewer than 181/226 tips")


if __name__ == "__main__":
    main()
