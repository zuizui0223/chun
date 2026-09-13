#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def header_only(path: Path) -> list[dict[str, object]]:
    wb = load_workbook(path, read_only=True, data_only=False)
    out = []
    for ws in wb.worksheets:
        header_row_index = None
        headers: list[str] = []
        # Read at most the first 12 rows and stop at the first row with >=2 nonempty cells.
        # Values below that row are never inspected by this schema-only step.
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True), start=1):
            nonempty = [v for v in row if v is not None and str(v).strip()]
            if len(nonempty) >= 2:
                header_row_index = i
                headers = [str(v).strip() if v is not None else '' for v in row]
                break
        out.append({
            'sheet': ws.title,
            'max_row': ws.max_row,
            'max_column': ws.max_column,
            'header_row_index': header_row_index,
            'headers': headers,
        })
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument('--input-dir', type=Path, required=True)
    ap.add_argument('--object-manifest', type=Path, required=True)
    ap.add_argument('--out', type=Path, required=True)
    a = ap.parse_args()

    xlsx = sorted(a.input_dir.glob('*.xlsx'))
    if not xlsx:
        raise SystemExit('no XLSX supplementary objects recovered')

    object_manifest = json.loads(a.object_manifest.read_text(encoding='utf-8'))
    result = {
        'version': 'v0.1',
        'status': 'SCHEMA_ONLY_NO_ROW_OUTCOME_INSPECTION',
        'source_pmcid': 'PMC7767864',
        'source_doi': '10.3389/fpls.2020.604389',
        'pmc_prefix': object_manifest['pmc_prefix'],
        'pmc_objects': object_manifest['objects'],
        'xlsx_files': [],
        'inspection_boundary': (
            'Only workbook metadata and the first plausible header row (within rows 1-12) were read. '
            'No cells below the detected header row were inspected or emitted.'
        ),
        'paper1_science_changed': False,
    }
    for path in xlsx:
        result['xlsx_files'].append({
            'filename': path.name,
            'bytes': path.stat().st_size,
            'sha256': sha256(path),
            'worksheets': header_only(path),
        })

    a.out.parent.mkdir(parents=True, exist_ok=True)
    a.out.write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
