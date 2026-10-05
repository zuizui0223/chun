#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import io
import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
STRUCT_PATH=ROOT/"scripts"/"preflight_ng2018_row3_structure_v0_1.py"
spec=importlib.util.spec_from_file_location("structure",STRUCT_PATH)
structure=importlib.util.module_from_spec(spec)
spec.loader.exec_module(structure)

HEADER_ROW=7
TRAIT_TERMS=("pelargonidin","cyanidin","delphinidin","anthocyanin","pigment","proportion","predominant")


def row_values(body:bytes,row:int)->list[str]:
    wb=load_workbook(io.BytesIO(body),read_only=True,data_only=False)
    ws=wb.worksheets[0]
    return ["" if c.value is None else str(c.value).strip() for c in ws[row]]


def source_unavailable_receipt(error: str) -> dict:
    return {
        "version": "v0.1",
        "status": "HOLD_NG2018_TRAIT_SUPPLEMENT_SOURCE_UNAVAILABLE",
        "source_error": error,
        "trait_source_url": None,
        "trait_source_sha256": None,
        "trait_source_bytes": None,
        "header_row_frozen_before_value_opening": HEADER_ROW,
        "header": [],
        "header_has_species": False,
        "trait_columns_matching_frozen_terms": [],
        "header_has_trait_column": False,
        "rows_8_plus_opened": False,
        "trait_data_rows_opened": 0,
        "trait_state_frequencies_computed": False,
        "tree_trait_crosswalk_computed": False,
        "hidden_memory_auc_computed": False,
        "next_gate": "STOP_HOLD",
        "interpretation_if_source_only": "Not evaluated because the exact frozen Table S1 source was unavailable.",
        "paper1_science_changed": False,
        "el_v0_3_science_changed": False,
        "v0_8_promotion_state_changed": False,
    }


def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--out",type=Path,required=True)
    a=ap.parse_args()
    a.out.parent.mkdir(parents=True,exist_ok=True)

    try:
        body,url=structure.preflight.recover_trait_source()
    except RuntimeError as exc:
        out=source_unavailable_receipt(str(exc))
        a.out.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
        print(json.dumps(out,indent=2,sort_keys=True))
        return 0

    structural=structure.sheet_structure(body)
    if structural["predeclared_header_candidate_row"] != HEADER_ROW:
        raise RuntimeError(
            f"frozen structural header candidate changed: "
            f"{structural['predeclared_header_candidate_row']} != {HEADER_ROW}"
        )

    header=row_values(body,HEADER_ROW)
    norm=[re.sub(r"\s+"," ",x.lower()).strip() for x in header]
    has_species=any(x=="species" or x.startswith("species ") for x in norm)
    trait_columns=[
        raw for raw,n in zip(header,norm)
        if any(term in n for term in TRAIT_TERMS)
    ]
    has_trait=bool(trait_columns)

    if not has_species:
        status="HOLD_NG2018_ROW7_NOT_SPECIES_HEADER"
    elif not has_trait:
        status="HOLD_NG2018_TABLE_S1_SOURCE_ONLY_NO_TRAIT_COLUMNS"
    else:
        status="NG2018_TRAIT_SCHEMA_PRESENT_CROSSWALK_GATE_NEXT"

    out={
        "version":"v0.1",
        "status":status,
        "trait_source_url":url,
        "trait_source_sha256":hashlib.sha256(body).hexdigest(),
        "trait_source_bytes":len(body),
        "header_row_frozen_before_value_opening":HEADER_ROW,
        "header":header,
        "header_has_species":has_species,
        "trait_columns_matching_frozen_terms":trait_columns,
        "header_has_trait_column":has_trait,
        "rows_8_plus_opened":False,
        "trait_data_rows_opened":0,
        "trait_state_frequencies_computed":False,
        "tree_trait_crosswalk_computed":False,
        "hidden_memory_auc_computed":False,
        "next_gate":(
            "FREEZE_IDENTIFIER_CROSSWALK"
            if status=="NG2018_TRAIT_SCHEMA_PRESENT_CROSSWALK_GATE_NEXT"
            else "STOP_HOLD"
        ),
        "interpretation_if_source_only":"Recovered Table S1 is a species/source inventory, not the row-level biochemical state matrix required by the frozen hidden-memory estimand.",
        "paper1_science_changed":False,
        "el_v0_3_science_changed":False,
        "v0_8_promotion_state_changed":False,
    }
    a.out.write_text(json.dumps(out,indent=2,sort_keys=True)+"\n")
    print(json.dumps(out,indent=2,sort_keys=True))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
