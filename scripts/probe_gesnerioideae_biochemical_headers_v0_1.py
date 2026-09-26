#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import urllib.request
import zipfile
from pathlib import Path

from openpyxl import load_workbook

EPMC="https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7767864/supplementaryFiles"
EXPECTED_XLSX_SHA256="a84abf67da0afb8c0bafd4c1251dbcbb6dc48eb286dca788e6e88ce0176ccbc8"
EXPECTED_SHEET="FINAL SAMPLE LIST"
HEADER_ROW=3
UA="CHUN-Gesnerioideae-biochem-header/0.1"


def sha256(b:bytes)->str:
    return hashlib.sha256(b).hexdigest()


def fetch(url:str)->bytes:
    req=urllib.request.Request(url,headers={"User-Agent":UA,"Accept":"*/*"})
    with urllib.request.urlopen(req,timeout=120) as r:
        return r.read()


def extract_header(xlsx:bytes,sheet:str,header_row:int)->list[str]:
    wb=load_workbook(io.BytesIO(xlsx),read_only=True,data_only=False)
    if sheet not in wb.sheetnames:
        raise ValueError(f"sheet missing: {sheet}")
    ws=wb[sheet]
    vals=[]
    for cell in ws[header_row]:
        v=cell.value
        vals.append("" if v is None else re.sub(r"\s+"," ",str(v).strip()))
    while vals and vals[-1]=="":
        vals.pop()
    return vals


def biochemical_column_candidates(header:list[str])->list[str]:
    out=[]
    for x in header:
        n=re.sub(r"\s+"," ",x.strip()).lower()
        if not n:
            continue
        if any(t in n for t in ("reflect","colour","color","hue","chittka","wavelength")):
            continue
        if any(t in n for t in (
            "anthocyan","anthocyanidin","pelargon","cyanidin","peonidin",
            "delphinidin","petunidin","malvidin","flavonoid","pigment",
            "hydroxyanth","deoxyanth","deo90","hyd90"
        )):
            out.append(x)
    return out


def locate_exact_workbook(container:bytes)->tuple[str,bytes]:
    with zipfile.ZipFile(io.BytesIO(container)) as z:
        matches=[]
        for name in z.namelist():
            if not name.lower().endswith(".xlsx"):
                continue
            b=z.read(name)
            if sha256(b)==EXPECTED_XLSX_SHA256:
                matches.append((name,b))
    if len(matches)!=1:
        raise RuntimeError(f"expected exactly one XLSX matching frozen SHA, got {len(matches)}")
    return matches[0]


def main()->int:
    ap=argparse.ArgumentParser()
    ap.add_argument("--out",type=Path,required=True)
    a=ap.parse_args()
    a.out.parent.mkdir(parents=True,exist_ok=True)

    container=fetch(EPMC)
    name,xlsx=locate_exact_workbook(container)
    header=extract_header(xlsx,EXPECTED_SHEET,HEADER_ROW)
    bio=biochemical_column_candidates(header)

    out={
      "version":"v0.1",
      "status":"GESNERIOIDEAE_BIOCHEMICAL_HEADER_GATE_READY",
      "source_doi":"10.3389/fpls.2020.604389",
      "pmcid":"PMC7767864",
      "selected_workbook":name,
      "workbook_sha256":sha256(xlsx),
      "sheet":EXPECTED_SHEET,
      "header_row":HEADER_ROW,
      "header":header,
      "biochemical_column_candidates":bio,
      "data_rows_read":0,
      "state_frequencies_computed":False,
      "hidden_memory_auc_computed":False,
      "old_three_level_schema_hold_changed":False,
      "next_gate":"FREEZE_TWO_LEVEL_BIOCHEMICAL_STATE_RULE_FROM_HEADERS_AND_SOURCE_DEFINITIONS",
      "paper1_science_changed":False,
      "el_v0_3_science_changed":False
    }
    a.out.write_text(json.dumps(out,indent=2,ensure_ascii=False,sort_keys=True)+"\n",encoding="utf-8")
    print(json.dumps(out,indent=2,ensure_ascii=False))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
