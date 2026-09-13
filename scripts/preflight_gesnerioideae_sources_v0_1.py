#!/usr/bin/env python3
import csv, hashlib, io, json, re, sys
from pathlib import Path
from urllib.parse import urljoin
import requests
from openpyxl import load_workbook
from Bio import Phylo

EPMC_XML = 'https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7767864/fullTextXML'
PMC_BIN = 'https://pmc.ncbi.nlm.nih.gov/articles/PMC7767864/bin/'
DRYAD_DATASET = 'https://datadryad.org/api/v2/datasets/doi:10.5061/dryad.m7589'
TREE_NAME = 'Gesne_Ago8_simple_combined_CA.tre'
OUT = Path('build/gesnerioideae_source_preflight_v0_1')


def sha256(b):
    return hashlib.sha256(b).hexdigest()


def get(url):
    r = requests.get(url, timeout=60)
    r.raise_for_status()
    return r


def norm_binomial(x):
    s = re.sub(r'[_]+', ' ', str(x).strip())
    s = re.sub(r'\s+', ' ', s)
    toks = s.split()
    if len(toks) < 2:
        return None
    return (toks[0].lower(), toks[1].lower())


def discover_supplementary_xlsx():
    xml = get(EPMC_XML).text
    hrefs = re.findall(r'(?:xlink:href|href)="([^"]+\.xlsx(?:\?[^\"]*)?)"', xml, flags=re.I)
    hrefs = list(dict.fromkeys(hrefs))
    if not hrefs:
        raise RuntimeError('No XLSX supplementary hrefs found in Europe PMC XML')
    rows = []
    for href in hrefs:
        if href.startswith('http'):
            url = href
        else:
            url = urljoin(PMC_BIN, href.split('/')[-1])
        b = get(url).content
        rows.append((href, url, b))
    return rows


def choose_s1_without_outcomes(xlsx_rows):
    candidates=[]
    for href,url,b in xlsx_rows:
        wb=load_workbook(io.BytesIO(b), read_only=True, data_only=False)
        for ws in wb.worksheets:
            header=[ws.cell(1,c).value for c in range(1, min(ws.max_column,80)+1)]
            header_s=[str(x).strip() if x is not None else '' for x in header]
            lower=[x.lower() for x in header_s]
            species_cols=[i+1 for i,x in enumerate(lower) if x in {'species','species name','species_name','taxon','taxon name'}]
            score=0
            if 170 <= ws.max_row <= 200: score += 3
            if species_cols: score += 5
            # header-only hints; no cell values beyond the species identifier column are inspected
            if any('voucher' in x for x in lower): score += 2
            if any('anthocyan' in x or 'pigment' in x for x in lower): score += 1
            if any('reflect' in x or 'color' in x or 'colour' in x for x in lower): score += 1
            candidates.append({'score':score,'href':href,'url':url,'bytes':b,'sheet':ws.title,'rows':ws.max_row,'cols':ws.max_column,'header':header_s,'species_cols':species_cols})
    candidates.sort(key=lambda x:(-x['score'], abs(x['rows']-181), x['href'], x['sheet']))
    if not candidates or not candidates[0]['species_cols']:
        raise RuntimeError('Could not identify Supplementary Table S1 from header/shape only')
    top=candidates[0]
    # require enough structural evidence to avoid result-oriented table selection
    if top['score'] < 8:
        raise RuntimeError(f'Ambiguous S1 structural identification: top score={top["score"]}')
    return top, candidates


def extract_species_only(s1):
    wb=load_workbook(io.BytesIO(s1['bytes']), read_only=True, data_only=False)
    ws=wb[s1['sheet']]
    c=s1['species_cols'][0]
    vals=[]
    for r in range(2,ws.max_row+1):
        v=ws.cell(r,c).value
        if v is not None and str(v).strip():
            vals.append(str(v).strip())
    keys=[norm_binomial(v) for v in vals]
    valid=[k for k in keys if k]
    return vals, valid


def download_dryad_tree():
    meta=get(DRYAD_DATASET).json()
    files_href=(meta.get('_links',{}).get('files',{}) or {}).get('href')
    if not files_href:
        # current Dryad v2 fallback
        ident=meta.get('identifier') or meta.get('id')
        if ident:
            files_href=f'https://datadryad.org/api/v2/versions/{ident}/files'
        else:
            raise RuntimeError('Dryad files endpoint absent')
    files=get(files_href).json()
    items=files.get('_embedded',{}).get('stash:files',[]) or files.get('files',[])
    match=None
    for it in items:
        if it.get('path')==TREE_NAME or it.get('filename')==TREE_NAME:
            match=it; break
    if not match:
        raise RuntimeError(f'{TREE_NAME} not found in Dryad file listing')
    dl=(match.get('_links',{}).get('stash:download',{}) or match.get('_links',{}).get('download',{})).get('href')
    if not dl:
        dl=match.get('downloadURL') or match.get('url')
    if not dl:
        raise RuntimeError('Dryad tree download URL absent')
    b=get(dl).content
    return meta, match, dl, b


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    xlsxs=discover_supplementary_xlsx()
    s1,cands=choose_s1_without_outcomes(xlsxs)
    species_rows, species_keys=extract_species_only(s1)
    meta, tree_item, tree_url, tree_b=download_dryad_tree()
    tree_path=OUT/TREE_NAME
    tree_path.write_bytes(tree_b)
    tree=Phylo.read(str(tree_path),'newick')
    tips=[t.name for t in tree.get_terminals() if t.name]
    tip_keys=[norm_binomial(x) for x in tips]
    tip_set={x for x in tip_keys if x}
    source_unique=[]; seen=set()
    for raw,key in zip(species_rows,species_keys):
        if key not in seen:
            seen.add(key); source_unique.append((raw,key))
    matched=[(raw,key) for raw,key in source_unique if key in tip_set]
    unmatched=[raw for raw,key in source_unique if key not in tip_set]
    receipt={
      'status':'GESNERIOIDEAE_SOURCE_PREFLIGHT_COMPLETE_OUTCOME_BLIND',
      'outcome_values_read':False,
      'auc_computed':False,
      'decision_computed':False,
      'supplementary_source':{
        'pmcid':'PMC7767864','selected_href':s1['href'],'selected_url':s1['url'],'sha256':sha256(s1['bytes']),
        'sheet':s1['sheet'],'max_row':s1['rows'],'max_column':s1['cols'],'header':s1['header'],
        'selection_basis':'workbook shape + header names only'
      },
      'species_identifier_only':{
        'source_rows_nonempty':len(species_rows),'unique_normalized_binomials':len(source_unique),
        'tree_tips':len(tips),'matched_unique_binomials':len(matched),'unmatched_source_species':unmatched
      },
      'tree_source':{
        'dryad_doi':'10.5061/dryad.m7589','filename':TREE_NAME,'download_url':tree_url,
        'sha256':sha256(tree_b),'tree_tips':len(tips)
      },
      'hard_stop': 'PASS_CROSSWALK_PREFLIGHT' if len(unmatched)==0 else 'HOLD_CROSSWALK',
      'forbidden_columns_read':False
    }
    (OUT/'source_preflight_receipt.json').write_text(json.dumps(receipt,indent=2,ensure_ascii=False)+'\n')
    (OUT/'species_only_crosswalk_preflight.csv').write_text('source_species,normalized_key,in_tree\n'+''.join(f'"{raw}","{key[0]} {key[1]}",{str(key in tip_set).lower()}\n' for raw,key in source_unique))
    print(json.dumps(receipt,indent=2,ensure_ascii=False))
    if receipt['hard_stop']!='PASS_CROSSWALK_PREFLIGHT':
        sys.exit(2)

if __name__=='__main__': main()
