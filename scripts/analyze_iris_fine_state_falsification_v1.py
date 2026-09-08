#!/usr/bin/env python3
"""Prospective matched falsification test for the surviving cross-radiation fine-state result.

IMPORTANT: the decision rule, state alphabet, coarse groupings, OpenTree matching,
permutation count, seed, and sensitivities are frozen in
analysis/iris_fine_state_falsification_prefreeze_v1.md before this endpoint is run.
"""
from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
import requests
from Bio import Phylo
from openpyxl import load_workbook

from inspect_iris_opentree_v1 import canonical_query, post, tnrs

SUPPLEMENT = "https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7588356/supplementaryFiles"
SOURCE_WORKBOOK_SHA256 = "183ef5231c48e782b0ea69a7aa5605d40c892dd91d9d9b2d259ed7b65d230072"
UA = "chun-iris-falsification/1.0"
OUT_DIR = Path("analysis/_generated/iris_falsification_v1")
SEED = 20260908
PERMUTATIONS = 9999
FINE = ("MAROON", "ORANGE", "PINK", "PURPLE", "RED", "YELLOW", "WHITE")
FINE_INDEX = {x: i for i, x in enumerate(FINE)}
CODE = {
    "mar": "MAROON",
    "ora": "ORANGE",
    "pin": "PINK",
    "pur": "PURPLE",
    "red": "RED",
    "yel": "YELLOW",
    "whi": "WHITE",
}
PRIMARY_COARSE = {
    "MAROON": "ANTHOCYANIN",
    "PINK": "ANTHOCYANIN",
    "PURPLE": "ANTHOCYANIN",
    "RED": "ANTHOCYANIN",
    "ORANGE": "CAROTENOID",
    "YELLOW": "CAROTENOID",
    "WHITE": "NO_MAJOR_PIGMENT",
}


def compact(x):
    return "" if x is None else re.sub(r"\s+", " ", str(x)).strip()


def load_traits():
    r = requests.get(SUPPLEMENT, headers={"User-Agent": UA}, timeout=90)
    r.raise_for_status()
    with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
        blob = zf.read("Table_1.xlsx")
    sha = hashlib.sha256(blob).hexdigest()
    if sha != SOURCE_WORKBOOK_SHA256:
        raise ValueError(f"publisher workbook SHA256 drift: {sha}")
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    ws = wb["Source of data"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [compact(x) for x in rows[0]]
    idx = {h: i for i, h in enumerate(headers) if h}
    if not {"Species", "Colour", "Pigment"} <= set(idx):
        raise ValueError("publisher workbook schema changed")

    traits = {}
    raw_pigment = {}
    for row in rows[1:]:
        species = compact(row[idx["Species"]])
        if not species:
            continue
        raw = compact(row[idx["Colour"]]).lower()
        tokens = [x.strip() for x in raw.split("&") if x.strip()]
        unknown = sorted(set(tokens) - set(CODE))
        if unknown:
            raise ValueError(f"unknown frozen Iris colour code for {species}: {unknown}")
        states = frozenset(FINE_INDEX[CODE[x]] for x in tokens)
        if not states:
            raise ValueError(f"empty colour vector for {species}")
        if species in traits:
            raise ValueError(f"duplicate source species: {species}")
        traits[species] = states
        raw_pigment[species] = compact(row[idx["Pigment"]]).lower()
    if len(traits) != 226:
        raise ValueError(f"expected 226 frozen source rows, got {len(traits)}")
    return traits, raw_pigment


def primary_coarse_key(states: frozenset[int]):
    return tuple(sorted({PRIMARY_COARSE[FINE[i]] for i in states}))


def white_nonwhite_key(states: frozenset[int]):
    return tuple(sorted({"WHITE" if FINE[i] == "WHITE" else "NONWHITE" for i in states}))


def build_tree_and_admission(traits):
    parsed = []
    for source in sorted(traits):
        query, normalization = canonical_query(source)
        parsed.append((source, query, normalization))
    canonical = [x[1] for x in parsed]
    matches = tnrs(canonical)

    rows = []
    by_ott = defaultdict(list)
    for (source, query, normalization), m in zip(parsed, matches):
        row = {
            "source_species": source,
            "query": query,
            "source_format_normalization": normalization,
            **{k: v for k, v in m.items() if k != "query"},
        }
        rows.append(row)
        if row["status"] == "EXACT" and row["ott_id"] is not None:
            by_ott[int(row["ott_id"])].append(row)

    collisions = {oid: rs for oid, rs in by_ott.items() if len(rs) > 1}
    collision_sources = {r["source_species"] for rs in collisions.values() for r in rs}
    for row in rows:
        if row["source_species"] in collision_sources:
            row["status"] = "REJECT_OTT_COLLISION"

    exact = [r for r in rows if r["status"] == "EXACT"]
    ott_to_source = {int(r["ott_id"]): r["source_species"] for r in exact}
    ids = sorted(ott_to_source)
    if not ids:
        return None, {}, rows, collisions, {"reason": "NO_EXACT_OTT_IDS"}

    sub = post("tree_of_life/induced_subtree", {"ott_ids": ids, "label_format": "id"})
    nwk = sub.get("newick", "")
    if not nwk:
        return None, {}, rows, collisions, {"reason": "EMPTY_INDUCED_SUBTREE", "synth_id": sub.get("synth_id", "")}
    tree = Phylo.read(io.StringIO(nwk), "newick")

    seen_ott = set()
    unexpected = []
    for tip in tree.get_terminals():
        label = str(tip.name).strip("'")
        m = re.search(r"(?:ott)?(\d+)$", label)
        if not m:
            unexpected.append(label)
            continue
        oid = int(m.group(1))
        if oid not in ott_to_source:
            unexpected.append(label)
            continue
        seen_ott.add(oid)
        tip.name = ott_to_source[oid]
    if unexpected:
        raise ValueError(f"unexpected OpenTree terminal labels: {unexpected[:10]}")

    tree_names = [x.name for x in tree.get_terminals()]
    if len(tree_names) != len(set(tree_names)):
        raise ValueError("duplicate source names after OpenTree relabelling")
    states = {name: traits[name] for name in tree_names}
    meta = {
        "synth_id": sub.get("synth_id", ""),
        "newick_sha256": hashlib.sha256(nwk.encode()).hexdigest(),
        "n_exact_tnrs_before_collisions": sum(r["status"] in {"EXACT", "REJECT_OTT_COLLISION"} for r in rows),
        "n_ott_collision_groups": len(collisions),
        "n_admitted_unique_ott": len(ids),
        "n_tree_overlap": len(tree_names),
        "coverage_fraction_of_226": len(tree_names) / 226.0,
        "missing_admitted_ott_ids_from_tree": sorted(set(ids) - seen_ott),
        "tnrs_status_counts": dict(Counter(r["status"] for r in rows)),
        "source_format_normalization_counts": dict(Counter(r["source_format_normalization"] for r in rows)),
    }
    return tree, states, rows, collisions, meta


def prune_tree(tree, keep):
    t = copy.deepcopy(tree)
    for tip in list(t.get_terminals()):
        if tip.name not in keep:
            t.prune(tip)
    names = {x.name for x in t.get_terminals()}
    if names != set(keep):
        raise ValueError("sensitivity pruning failed")
    return t


def sankoff_conditional(tree, states, key_fn, label):
    leaves = list(tree.get_terminals())
    leaf_names = [x.name for x in leaves]
    if len(leaf_names) != len(set(leaf_names)) or set(leaf_names) != set(states):
        raise ValueError(f"{label}: tree/state one-to-one mismatch")

    nodes = list(tree.find_clades(order="postorder"))
    node_index = {id(node): i for i, node in enumerate(nodes)}
    children = [[node_index[id(child)] for child in node.clades] for node in nodes]
    leaf_nodes = [node_index[id(x)] for x in leaves]
    nt = len(leaves)
    k = len(FINE)
    big = np.int16(10000)

    patterns = np.full((nt, k), big, dtype=np.int16)
    for pos, name in enumerate(leaf_names):
        for s in states[name]:
            patterns[pos, s] = 0

    by_group = defaultdict(list)
    for pos, name in enumerate(leaf_names):
        by_group[key_fn(states[name])].append(pos)
    groups = [np.asarray(v, dtype=int) for _, v in sorted(by_group.items(), key=lambda kv: str(kv[0]))]

    rng = np.random.default_rng(SEED)
    all_scores = []
    for start in range(0, PERMUTATIONS + 1, 500):
        n = min(500, PERMUTATIONS + 1 - start)
        indices = np.tile(np.arange(nt, dtype=int), (n, 1))
        for j in range(n):
            if start + j == 0:
                continue
            for members in groups:
                indices[j, members] = rng.permutation(members)

        d = np.zeros((len(nodes), n, k), dtype=np.int16)
        for pos, node_i in enumerate(leaf_nodes):
            d[node_i] = patterns[indices[:, pos]]
        for i in range(len(nodes)):
            if not children[i]:
                continue
            d[i].fill(0)
            for c in children[i]:
                child = d[c]
                minimum = child.min(axis=1)
                d[i] += np.minimum(child, minimum[:, None] + 1)
        all_scores.extend(d[-1].min(axis=1).astype(int).tolist())

    observed = int(all_scores[0])
    null = np.asarray(all_scores[1:], dtype=float)
    null_mean = float(null.mean())
    ratio = float(observed / null_mean) if null_mean else None
    p_lower = float((1 + int((null <= observed).sum())) / (len(null) + 1))
    passed = bool(p_lower <= 0.01 and ratio is not None and ratio < 1)
    return {
        "label": label,
        "n_tips": nt,
        "fine_state_tip_presence": {
            FINE[i]: int(sum(i in states[name] for name in leaf_names)) for i in range(k)
        },
        "conditional_group_sizes": {"|".join(map(str, key)): len(v) for key, v in sorted(by_group.items(), key=lambda kv: str(kv[0]))},
        "observed_minimum_changes": observed,
        "null_mean": null_mean,
        "null_quantiles": np.quantile(null, [0, 0.025, 0.5, 0.975, 1]).tolist(),
        "observed_over_null_mean": ratio,
        "p_lower": p_lower,
        "permutations": PERMUTATIONS,
        "seed": SEED,
        "support_gate_p_le_0_01_and_ratio_lt_1": passed,
    }


def admission_gate(states, tree_meta):
    coarse_single = defaultdict(list)
    for name, st in states.items():
        key = primary_coarse_key(st)
        if len(key) == 1:
            coarse_single[key[0]].append(name)
    class_details = {}
    biological_class_gate = False
    for key, names in sorted(coarse_single.items()):
        fine_union = sorted({FINE[i] for name in names for i in states[name]})
        class_details[key] = {"n_tips": len(names), "distinct_fine_states": fine_union}
        if len(names) >= 30 and len(fine_union) >= 3:
            biological_class_gate = True
    gates = {
        "coverage_ge_181_of_226": tree_meta.get("n_tree_overlap", 0) >= 181,
        "all_admitted_tips_have_fine_state": bool(states) and all(states.values()),
        "coarse_class_ge_30_tips_and_ge_3_fine_states": biological_class_gate,
        "no_signal_based_taxon_exclusion": True,
    }
    return gates, class_details, all(gates.values())


def classify(primary, sensitivities):
    trait = [sensitivities["single_coarse_only"], sensitivities["white_nonwhite"]]
    available = [x for x in trait if x.get("status") == "COMPUTED"]
    if len(available) < 2:
        if primary["support_gate_p_le_0_01_and_ratio_lt_1"] and all(x["result"]["support_gate_p_le_0_01_and_ratio_lt_1"] for x in available):
            return "SUPPORTIVE_ALIGNMENT"
        return "ADVERSE_BUT_NOT_REFUTATION"
    primary_pass = primary["support_gate_p_le_0_01_and_ratio_lt_1"]
    sensitivity_passes = [x["result"]["support_gate_p_le_0_01_and_ratio_lt_1"] for x in available]
    if primary_pass and all(sensitivity_passes):
        return "SUPPORTIVE_ALIGNMENT"
    if (not primary_pass) and not any(sensitivity_passes):
        return "REFUTATION"
    return "MIXED"


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    traits, raw_pigment = load_traits()
    tree, states, tnrs_rows, collisions, tree_meta = build_tree_and_admission(traits)
    gates, class_details, admitted = admission_gate(states, tree_meta)

    base = {
        "version": "v1",
        "test_unit": "Iris",
        "source_doi": "10.3389/fpls.2020.569811",
        "source_workbook_sha256": SOURCE_WORKBOOK_SHA256,
        "frozen_source_rows": 226,
        "fine_alphabet": list(FINE),
        "primary_coarse_mapping": PRIMARY_COARSE,
        "permutations": PERMUTATIONS,
        "seed": SEED,
        "tree": tree_meta,
        "admission_gates": gates,
        "coarse_class_admission_details": class_details,
        "admission_pass": admitted,
        "topology_sensitivity": "NOT_AVAILABLE_SOURCE_SUPPLEMENT_NO_MACHINE_READABLE_FINAL_NEWICK",
        "claim_boundary": "Prospective fourth-radiation matched test of residual fine-state phylogenetic organization after coarse conditioning; not a universal transition-direction or common-mechanism test.",
        "paper1_science_changed": False,
    }

    if not admitted or tree is None:
        result = {
            **base,
            "classification": "HOLD",
            "primary": None,
            "sensitivities": {},
            "interpretation": "Frozen OpenTree/observation admission gate failed before the colour-signal endpoint; no biological counterexample is claimed.",
        }
    else:
        primary = sankoff_conditional(tree, states, primary_coarse_key, "PRIMARY_SOURCE_PIGMENT_GROUPS")

        single_names = {name for name, st in states.items() if len(primary_coarse_key(st)) == 1}
        if len(single_names) >= 3:
            single_tree = prune_tree(tree, single_names)
            single_states = {name: states[name] for name in single_names}
            single = {
                "status": "COMPUTED",
                "result": sankoff_conditional(single_tree, single_states, primary_coarse_key, "SENSITIVITY_SINGLE_COARSE_ONLY"),
            }
        else:
            single = {"status": "NOT_ADMISSIBLE", "reason": "fewer than 3 single-coarse tips"}

        white_nonwhite = {
            "status": "COMPUTED",
            "result": sankoff_conditional(tree, states, white_nonwhite_key, "SENSITIVITY_WHITE_NONWHITE"),
        }
        sensitivities = {
            "single_coarse_only": single,
            "white_nonwhite": white_nonwhite,
            "source_topology": {"status": "NOT_AVAILABLE"},
        }
        classification = classify(primary, sensitivities)
        if classification == "SUPPORTIVE_ALIGNMENT":
            interpretation = "Iris prospectively replicates residual fine-state phylogenetic organization after coarse conditioning; the cross-radiation recurrence extends from 3/3 to 4/4 matched testable radiations under this endpoint."
        elif classification == "REFUTATION":
            interpretation = "Iris is a matched prospective fourth-radiation failure under the frozen primary and trait-sensitivity tests; the prior 3/3 recurrence must not be generalized as a universal cross-radiation fine-state rule."
        elif classification == "MIXED":
            interpretation = "Iris gives materially different primary and frozen-sensitivity outcomes; the surviving recurrence claim is unresolved rather than cleanly replicated or refuted in this fourth radiation."
        else:
            interpretation = "Iris is adverse in the primary endpoint but robustness is incomplete; retain adverse-but-not-refutation status."
        result = {
            **base,
            "classification": classification,
            "primary": primary,
            "sensitivities": sensitivities,
            "interpretation": interpretation,
        }

    (OUT_DIR / "result.json").write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    audit = {
        "tnrs_rows": tnrs_rows,
        "ott_collision_groups": {
            str(oid): [{k: r[k] for k in ("source_species", "query", "matched_name", "ott_id")} for r in rs]
            for oid, rs in collisions.items()
        },
        "raw_pigment_qc_counts": dict(Counter(raw_pigment.values())),
    }
    (OUT_DIR / "admission_audit.json").write_text(json.dumps(audit, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("IRIS_FALSIFICATION_RESULT=" + json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
